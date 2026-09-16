from typing import List, Dict, Any
import docx

def load_docx(file_path: str) -> List[Dict[str, Any]]:
    """
    Extract text from a DOCX file using python-docx.
    Tracks headings as sections.
    """
    doc = docx.Document(file_path)
    extracted_items = []
    
    current_section = "General"
    current_paragraphs = []
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
            
        # Check if paragraph has heading style
        if para.style and para.style.name and 'Heading' in para.style.name:
            # If we had accumulated paragraphs under previous section, flush them
            if current_paragraphs:
                extracted_items.append({
                    "text": "\n".join(current_paragraphs),
                    "page_number": None,
                    "section": current_section
                })
                current_paragraphs = []
            current_section = text
        else:
            current_paragraphs.append(text)
            
    # Also extract any tables in docx
    for table_idx, table in enumerate(doc.tables):
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                table_rows.append(" | ".join(row_cells))
        if table_rows:
            current_paragraphs.append(f"[Table {table_idx + 1}]:\n" + "\n".join(table_rows))
            
    if current_paragraphs:
        extracted_items.append({
            "text": "\n".join(current_paragraphs),
            "page_number": None,
            "section": current_section
        })
        
    return extracted_items
